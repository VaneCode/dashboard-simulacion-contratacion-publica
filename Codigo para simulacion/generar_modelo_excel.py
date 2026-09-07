"""Genera el libro Excel de resultados a partir de los JSON del modelo.

Ejecutar después de: python generate_tesis_data.py
"""
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "generated_data"
OUT = ROOT / "Modelo_PAC_POA_Contratacion_2026_reproducido.xlsx"
NAVY, BLUE, WHITE, YELLOW = "17365D", "2F75B5", "FFFFFF", "FFF2CC"

def read(name):
    return json.loads((DATA / f"{name}.json").read_text(encoding="utf-8"))

def title(ws, text, columns):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=columns)
    c = ws.cell(1, 1, text)
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.font = Font(color=WHITE, bold=True, size=16)
    c.alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 25

def add_table(ws, start_row, headers, rows, name):
    for col, value in enumerate(headers, 1):
        c = ws.cell(start_row, col, value)
        c.fill = PatternFill("solid", fgColor=BLUE)
        c.font = Font(color=WHITE, bold=True)
    for row_i, row in enumerate(rows, start_row + 1):
        for col_i, value in enumerate(row, 1):
            ws.cell(row_i, col_i, value)
    end_row = start_row + len(rows)
    end_col = get_column_letter(len(headers))
    tab = Table(displayName=name, ref=f"A{start_row}:{end_col}{end_row}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tab)
    ws.freeze_panes = f"A{start_row + 1}"

def widths(ws, values):
    for index, width in enumerate(values, 1):
        ws.column_dimensions[get_column_letter(index)].width = width

pac, poa = read("pac_limpio"), read("poa_sintetico")
cycle, reps = read("ciclo_sintetico"), read("sensibilidad_replicaciones")
sens, params, dictionary, dash = [read(x) for x in ("sensibilidad_resumen", "parametros", "diccionario", "dashboard")]
wb = Workbook(); wb.remove(wb.active)

ws = wb.create_sheet("Dashboard"); title(ws, "Modelo integrado PAC–POA 2026", 8)
cards = [("Registros PAC", dash["registros_pac"]), ("Monto PAC", dash["monto_planificado"]),
         ("Objetivos POA", dash["objetivos_poa"]), ("Materialización", dash["metricas_base"]["tasa_materializacion_pct"] / 100),
         ("Finalización", dash["metricas_base"]["tasa_finalizacion_pct"] / 100),
         ("Días a adjudicar", dash["metricas_base"]["mediana_dias_adjudicacion"]),
         ("Desiertos/cancelados", dash["metricas_base"]["desiertos_cancelados"])]
add_table(ws, 3, ["Indicador", "Resultado"], cards, "tblDashboard")
ws["A13"] = "Escenario"; ws["B13"] = "Materialización p05"; ws["C13"] = "Mediana"; ws["D13"] = "p95"; ws["E13"] = "Finalización mediana"; ws["F13"] = "Días"
for c in ws[13]: c.fill=PatternFill("solid",fgColor=NAVY); c.font=Font(color=WHITE,bold=True)
for i, x in enumerate(sens, 14):
    ws.append([x["escenario"], x["tasa_materializacion_pct_p05"]/100, x["tasa_materializacion_pct_mediana"]/100,
               x["tasa_materializacion_pct_p95"]/100, x["tasa_finalizacion_pct_mediana"]/100, x["mediana_dias_adjudicacion_mediana"]])
ws.merge_cells("A20:H22"); ws["A20"]="ADVERTENCIA: POA, fechas, estados y montos de adjudicación/ejecución son sintéticos. No representan desempeño real del GAD Municipal de Guayaquil."
ws["A20"].fill=PatternFill("solid",fgColor=YELLOW); ws["A20"].alignment=Alignment(wrap_text=True,vertical="center")
widths(ws,[25,20,18,18,22,14,14,14])

ws=wb.create_sheet("PAC_Limpio")
headers=list(pac[0].keys()); title(ws,"PAC 2026 depurado — fuente observada y campos derivados",len(headers))
add_table(ws,3,headers,[[r.get(h) for h in headers] for r in pac],"tblPAC")
widths(ws,[18,10,18,16,16,18,60,12,15,16,18,14,20,18,42,14,20,24,20])

ws=wb.create_sheet("POA_Sintetico"); headers=list(poa[0].keys()); title(ws,"POA sintético vinculado al PAC",len(headers))
add_table(ws,3,headers,[[r.get(h) for h in headers] for r in poa],"tblPOA"); widths(ws,[18,18,48,38,48,18,22,22,15,40])

ws=wb.create_sheet("Ciclo_Sintetico"); headers=list(cycle[0].keys()); title(ws,"Ciclo contractual sintético — escenario base, semilla 2026",len(headers))
add_table(ws,3,headers,[[r.get(h) for h in headers] for r in cycle],"tblCiclo"); widths(ws,[20,20,20,15,42,20,18,18,18,18,18,20,24,24,22,40])

ws=wb.create_sheet("Sensibilidad"); headers=list(reps[0].keys()); title(ws,"Análisis de sensibilidad — 100 réplicas por escenario",len(headers))
add_table(ws,3,headers,[[r.get(h) for h in headers] for r in reps],"tblSensibilidad"); widths(ws,[18,12,14,22,20,24,20,20,20,22,18,14])

ws=wb.create_sheet("Parametros"); headers=list(params[0].keys()); title(ws,"Parámetros de simulación",len(headers))
add_table(ws,3,headers,[[r.get(h) for h in headers] for r in params],"tblParametros"); widths(ws,[50,22,28])

ws=wb.create_sheet("Diccionario"); headers=list(dictionary[0].keys()); title(ws,"Diccionario de datos",len(headers))
add_table(ws,3,headers,[[r.get(h) for h in headers] for r in dictionary],"tblDiccionario"); widths(ws,[35,18,85])

for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row: c.alignment = Alignment(vertical="top", wrap_text=True)
wb.save(OUT)
print(f"Archivo generado: {OUT}")
